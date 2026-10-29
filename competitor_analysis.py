"""任意アップロードの競合還元額調査分析。予測ロジックには依存しない。"""
import datetime as dt
import math
import re
from io import BytesIO

import plotly.graph_objects as go
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

DATE_SHEET = re.compile(r"^\d{6}$")
AMOUNT = re.compile(r"[-+]?\d[\d,]*(?:\.\d+)?")


@st.cache_data(show_spinner=False, max_entries=3)
def parse_competitor_workbook(content: bytes):
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    records = []
    warnings = []
    for name in wb.sheetnames:
        if not DATE_SHEET.fullmatch(name):
            continue
        try:
            date = dt.datetime.strptime(name, "%y%m%d").date()
        except ValueError:
            warnings.append(f"日付として認識できないシート: {name}")
            continue
        ws = wb[name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or str(header[0]).strip() != "媒体名":
            warnings.append(f"見出し形式が異なるシート: {name}")
            continue
        products = [(i, str(v).strip()) for i, v in enumerate(header) if i >= 3 and v is not None and str(v).strip()]
        for row in rows:
            if not row or row[0] is None:
                continue
            media = str(row[0]).strip()
            segment = str(row[1]).strip() if len(row) > 1 and row[1] is not None else "未設定"
            category = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            if category != "単価":
                continue
            for idx, product in products:
                if idx >= len(row) or row[idx] is None:
                    continue
                raw = str(row[idx]).strip()
                match = AMOUNT.search(raw.replace("，", ","))
                if not match:
                    continue
                value = float(match.group().replace(",", ""))
                if not math.isfinite(value) or value < 0:
                    continue
                records.append((date, media, segment, product, value))
    wb.close()
    df = pd.DataFrame(records, columns=["調査日", "媒体", "セグメント", "商品", "還元額"])
    if not df.empty:
        df["調査日"] = pd.to_datetime(df["調査日"])
        df = df.drop_duplicates(["調査日", "媒体", "セグメント", "商品"], keep="last")
    return df, warnings


def show_competitor_analysis():
    with st.expander("📈 競合還元額分析（任意アップロード）", expanded=False):
        st.info("競合還元額調査ファイルは [こちらのBox格納先](https://rak.box.com/s/uaf1thskosdl8eif8wb0cjiksd59ci81) からダウンロードできます。アップロードは任意です。")
        file = st.file_uploader("競合還元額調査Excel", type=["xlsx", "xlsm"], key="competitor_workbook", help="シート名がYYMMDDの日付シートを読み取ります。未アップロードでも既存機能はそのまま利用できます。")
        if file is None:
            st.info("調査Excelをアップロードすると、媒体別の最新還元額と推移を表示します。")
            return
        try:
            with st.spinner("日付シートを読み込み中…"):
                df, warnings = parse_competitor_workbook(file.getvalue())
        except Exception as exc:
            st.error(f"競合調査ファイルを読み取れませんでした: {exc}")
            return
        if df.empty:
            st.warning("対象の日付シートから単価データを取得できませんでした。")
            return
        if warnings:
            st.warning(f"読み取り対象外のシートが{len(warnings)}件あります。例: {', '.join(warnings[:3])}")
        st.caption(f"収録調査日: {df['調査日'].min():%Y/%m/%d} ～ {df['調査日'].max():%Y/%m/%d} / {df['調査日'].nunique():,}日")
        preferred = ["楽天カード", "Paypayカード", "三井住友カード（NL）",
                     "エポスカード", "JCB CARD W", "イオンカード"]
        product_options = sorted(df["商品"].unique())
        defaults = [p for p in preferred if p in product_options]
        c1, c2 = st.columns([2, 1])
        with c1:
            products = st.multiselect(
                "対象商品（複数選択可）", product_options, default=defaults,
                key="competitor_products",
            )
        with c2:
            period = st.selectbox(
                "表示期間", [30, 90, 180, 365, "全期間"], index=0,
                format_func=lambda x: f"過去{x}日" if isinstance(x, int) else x,
                key="competitor_period",
            )
        segment_options = sorted(df["セグメント"].unique())
        segments = st.multiselect(
            "セグメント", segment_options, default=segment_options,
            key="competitor_segments",
        )
        if not products or not segments:
            st.info("商品とセグメントを選択してください。")
            return
        base = df[df["商品"].isin(products) & df["セグメント"].isin(segments)].copy()
        if base.empty:
            st.info("選択条件に該当するデータがありません。")
            return
        end = base["調査日"].max()
        if period != "全期間":
            base = base[base["調査日"] >= end - pd.Timedelta(days=period - 1)]
        options = sorted(base["媒体"].unique())
        media = st.multiselect(
            "表示媒体", options, default=options[:min(8, len(options))],
            key="competitor_media",
        )
        selected = base[base["媒体"].isin(media)].copy()
        if selected.empty:
            st.info("表示する媒体を選択してください。")
            return
        latest_date = selected["調査日"].max()
        latest = selected[selected["調査日"] == latest_date].sort_values("還元額", ascending=False)
        k1, k2, k3 = st.columns(3)
        k1.metric("最新調査日", latest_date.strftime("%Y/%m/%d"))
        k2.metric("最新最高還元額", f"{latest['還元額'].max():,.0f}円")
        k3.metric("期間内最高還元額", f"{selected['還元額'].max():,.0f}円")
        st.subheader("媒体別・商品別還元額推移")
        st.caption("横軸＝調査日、縦軸＝還元額。媒体名を左に、推移グラフを右に並べます。商品ごとに色分けし、セグメントは別線にします。")
        low, high = selected["還元額"].min(), selected["還元額"].max()
        lower = max(0, math.floor((low - 500) / 1000) * 1000)
        upper = math.ceil((high + 500) / 1000) * 1000
        if upper <= lower:
            upper = lower + 1000
         for media_name in media:
            part = selected[selected["媒体"] == media_name]
            if part.empty:
                continue
            label_col, graph_col = st.columns([1, 5], vertical_alignment="center")
            with label_col:
                st.markdown(f"**{media_name}**")
            with graph_col:
                fig = go.Figure()
                for (product_name, segment_name), group in part.groupby(["商品", "セグメント"], sort=True):
                    series = group.sort_values("調査日").drop_duplicates("調査日", keep="last")
                    label = product_name if len(segments) == 1 else f"{product_name} / {segment_name}"
                    fig.add_trace(go.Scatter(
                        x=series["調査日"], y=series["還元額"],
                        mode="lines", line_shape="hv", name=label,
                        hovertemplate="%{x|%Y/%m/%d}<br>%{y:,.0f}円<extra>%{fullData.name}</extra>",
                        connectgaps=False,
                    ))
                fig.update_layout(
                    height=260, margin=dict(l=10, r=10, t=15, b=20),
                    legend=dict(orientation="h", y=-0.32, x=0),
                    hovermode="x unified",
                    xaxis=dict(
                        title=None,
                        range=[selected["調査日"].min(), selected["調査日"].max() + pd.Timedelta(days=1)],
                        tickformat="%m/%d",
                        dtick=86400000,  # 1日 = 86,400,000ミリ秒
                        tickangle=-55,
                        tickfont=dict(size=9),
                        showgrid=False,
                    ),
                    yaxis=dict(
                        title=None, range=[lower, upper], tick0=0, dtick=1000,
                        tickformat=",d", ticksuffix="円", showgrid=False,
                    ),
                )
                st.plotly_chart(fig, use_container_width=True, key=f"competitor_chart_{media.index(media_name)}")
        st.caption("縦軸は1,000円刻み・全媒体共通の自動範囲。還元額は丸めず、調査日間を階段表示します。")
        st.subheader("最新調査日の還元額一覧")
        st.dataframe(
            latest[["媒体", "商品", "セグメント", "還元額"]],
            hide_index=True, width="stretch",
            column_config={"還元額": st.column_config.NumberColumn(format="%.0f 円")},
        )
        st.caption("0円は元ファイル記載値のまま表示します。掲載なしとの区別は元データ側の定義確認が必要です。")
